# -*- coding: utf-8 -*-
"""Slovak user guide as a Word file, with both boards' UI screenshots."""

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
IMG128 = ROOT / "docs" / "img" / "ws128"
IMG35 = ROOT / "docs" / "img" / "ws35"
OUT = ROOT / "docs" / "GlukoView-navod.docx"


def on_black(path: Path) -> io.BytesIO:
    im = Image.open(path).convert("RGBA")
    bg = Image.new("RGB", im.size, (0, 0, 0))
    bg.paste(im, mask=im.split()[-1])
    buf = io.BytesIO()
    bg.save(buf, format="PNG")
    buf.seek(0)
    return buf


def picture_source(path: Path):
    if "ws128" in path.parts:
        return on_black(path)
    return str(path)


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="C4A0A0", sz="8"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def no_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def add_text(doc, text, *, size=11, bold=False, space_after=8, space_before=0, center=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.08
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0])
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_heading_sk(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=18 if level == 1 else 14, bold=True, color=RGBColor(0x14, 0x3D, 0x4A))
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p


def caption_run(paragraph, text):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, italic=True, color=RGBColor(0x55, 0x55, 0x55))


def add_picture(doc, path: Path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(picture_source(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    caption_run(cap, caption)


def add_pair(doc, items):
    """items: list of (path, width_cm, caption)."""
    table = doc.add_table(rows=1, cols=len(items))
    table.autofit = True
    no_table_borders(table)
    for cell, (path, width_cm, caption) in zip(table.rows[0].cells, items):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        run.add_picture(picture_source(path), width=Cm(width_cm))
        cap = cell.add_paragraph()
        caption_run(cap, caption)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    for folder in (IMG128, IMG35):
        if not folder.is_dir():
            raise SystemExit(f"chýbajú obrázky: {folder}")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("GlukoView  ·  nie je zdravotnícka pomôcka  ·  obrázky sú ukážka, firmvér 2.85")
    set_run_font(fr, size=8, color=RGBColor(0x88, 0x88, 0x88))

    add_text(doc, "GlukoView", size=26, bold=True, space_after=0, color=RGBColor(0x14, 0x3D, 0x4A))
    add_text(doc, "Návod na použitie", size=16, space_after=10, color=RGBColor(0x3A, 0x7A, 0x4A))

    warn = doc.add_table(rows=1, cols=1)
    cell = warn.cell(0, 0)
    shade_cell(cell, "FDECEC")
    set_cell_border(cell, "C45050", "12")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Toto nie je zdravotnícka pomôcka.")
    set_run_font(r, size=12, bold=True, color=RGBColor(0x8C, 0x1D, 0x1D))
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(
        "Hodnota na displeji je len orientačná. Nesmie byť jediný podklad na dávku inzulínu "
        "ani na rozhodnutie o liečbe. Over si ju na senzore alebo na čítačke, ktorú používaš s lekárom."
    )
    set_run_font(r2, size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading_sk(doc, "Čo to je", 1)
    add_text(
        doc,
        "GlukoView je stolný ukazovateľ glykémie. Nie sú to hodinky. Na obrazovke ukáže posledné "
        "číslo, ktoré stiahne z nastaveného zdroja, spolu so šípkou trendu a vekom merania. "
        "Menu na displeji je po slovensky. V nastavení sa dá prepnúť aj čeština, angličtina, "
        "nemčina, španielčina a francúzština.",
    )
    add_text(
        doc,
        "Sú dve dosky. Obraz z jednej sa nesmie nahrať do druhej: displej, piny aj pamäť sú iné.",
    )
    add_bullet(doc, "Kruhový Waveshare 1,28″ (GC9A01). Obrazovka je kruh 240×240. Nemá zvuk.")
    add_bullet(doc, "Široký Waveshare ESP32-S3-Touch-LCD-3.5B. Obrazovka je na šírku 480×320. Má reproduktor a vie zobraziť dve osoby vedľa seba.")
    add_text(
        doc,
        "Čísla na obrázkoch sú vymyslené, aby bolo vidno farby a menu. Nie je to živé meranie. "
        "Meno Anna a Adam je len ukážka, ako široká doska pomenuje dve osoby.",
        space_before=6,
    )

    add_pair(
        doc,
        [
            (IMG128 / "face-ok.png", 6.2, "Kruh, jedna osoba. 6,4 mmol, v cieli, meranie staré 3 minúty."),
            (IMG35 / "face-ok.png", 10.2, "Široká doska, režim jednej osoby. Meno je vľavo, keď sú nastavené dve osoby."),
        ],
    )

    add_heading_sk(doc, "Prvé zapnutie", 1)
    add_text(
        doc,
        "Pri prvom štarte sa pýta na jazyk. Ťukni na ten, ktorý chceš. Potom sa na chvíľu ukáže logo a verzia firmvéru.",
    )
    add_pair(
        doc,
        [
            (IMG128 / "lang.png", 5.6, "Výber jazyka na kruhu."),
            (IMG35 / "lang.png", 10.6, "Ten istý výber na širokej doske."),
        ],
    )
    add_pair(
        doc,
        [
            (IMG128 / "splash.png", 5.6, "Logo po štarte, kruh."),
            (IMG35 / "splash.png", 10.6, "Logo po štarte, široká doska. Číslo dole je verzia."),
        ],
    )
    add_text(doc, "Ak doska ešte nepozná domácu sieť, na displeji sa striedajú tieto dva stavy.")
    add_pair(
        doc,
        [
            (IMG35 / "boot-connect.png", 8.2, "Pripája sa na uloženú sieť."),
            (IMG35 / "boot-setup.png", 8.2, "Keď sieť ešte nie je, postup je na displeji."),
        ],
    )
    add_text(doc, "Postup z obrazovky, prvýkrát:", space_after=4)
    add_bullet(doc, "Na telefóne alebo počítači sa pripoj na WiFi s názvom GlukoView. Toto prvé WiFi heslo nemá. Po pripojení na domácu sieť tento bod zhasne.")
    add_bullet(doc, "V prehliadači otvor adresu 10.0.0.1.")
    add_bullet(doc, "Vyber domácu sieť a ulož ju. Doska sa na ňu pripojí a prístupový bod GlukoView skončí.")
    add_bullet(doc, "Stránka potom ostane na IP adrese dosky v domácej sieti. Tam sa dopĺňajú účty zdroja. Adresu neskôr uvidíš v Informáciách na displeji.")
    add_bullet(doc, "Nastav štvormiestny PIN. Kým PIN nie je, stránka sa otvorí, ale heslá k účtom do nej nevypisuje.")
    add_text(
        doc,
        "Krátke stlačenie tlačidla BOOT (kratšie ako dve sekundy) otvorí menu. To isté spraví podržanie prsta na obrazovke. "
        "Ak BOOT podržíš už pri zapnutí asi päť sekúnd, doska zmaže WiFi aj nastavenia a reštartuje sa.",
    )

    add_heading_sk(doc, "Ako čítať číslo", 1)
    add_text(
        doc,
        "Veľké číslo je glykémia. Predvolene je v mmol/l. V menu Vzhľad sa dajú prepnúť jednotky na mg/dl. "
        "Šípka je trend zo zdroja. Malé číslo vedľa šípky je zmena od predchádzajúcej hodnoty: "
        "oranžová, keď stúpa, modrá, keď klesá, sivá, keď sa takmer nehla. Dá sa vypnúť. "
        "Pod tým je vek merania v minútach.",
    )
    add_text(doc, "Farba kruhu pri predvolených limitoch 3,9 / 10,0 / 15,0 mmol/l:", space_after=4)
    add_bullet(doc, "Zelená, keď je hodnota od spodného limitu po stred pásma.")
    add_bullet(doc, "Zelená prechádza do žltej, keď sa v pásme blíži k vrchnému limitu.")
    add_bullet(doc, "Oranžová, keď je nad vrchným limitom a ešte pod kritickou hodnotou.")
    add_bullet(doc, "Červená, keď je na spodnom limite alebo pod ním, a tiež keď dosiahne kritickú hodnotu alebo ju prekročí.")
    add_bullet(doc, "Sivý kruh a nápis Staré, keď je meranie staršie ako 15 minút. Číslo vtedy ber ako neaktuálne a over ho na senzore.")

    add_pair(
        doc,
        [
            (IMG128 / "face-ok.png", 6.4, "V cieli. Šípka mierne hore, zmena +0,3, vek 3 min."),
            (IMG128 / "face-low.png", 6.4, "Pod spodným limitom. Kruh je červený."),
        ],
    )
    add_pair(
        doc,
        [
            (IMG128 / "face-high.png", 6.4, "Nad vrchným limitom, ešte nie kritické. Kruh je oranžový."),
            (IMG128 / "face-critical.png", 6.4, "Kritická hodnota. Kruh je červený."),
        ],
    )
    add_picture(
        doc,
        IMG128 / "face-stale.png",
        6.4,
        "Meranie staršie ako 15 minút. Nápis Staré a vek sú červené, kruh je sivý.",
    )
    add_text(
        doc,
        "Šípky, ktoré zdroj posiela: ↑↑ veľmi rýchlo hore, ↑ hore, ↗ mierne hore, → rovno, "
        "↘ mierne dole, ↓ dole, ↓↓ veľmi rýchlo dole. Nie každý zdroj posiela všetky. "
        "Libre používa kratšiu sadu bez dvojitých šípok.",
    )

    add_heading_sk(doc, "Ďalšie vzhľady ciferníka", 1)
    add_text(
        doc,
        "Potiahnutie doľava prejde na ďalší vzhľad, potiahnutie doprava sa vráti. "
        "Pri jednej osobe sú tri: veľké číslo, hodiny a graf. "
        "Potiahnutie hore alebo dole požiada o novú hodnotu hneď, nečaká na ďalšie automatické stiahnutie.",
    )
    add_pair(
        doc,
        [
            (IMG128 / "face-clock.png", 6.4, "Hodiny. Čas je miestny čas dosky. 09:41 na obrázku je len ukážka."),
            (IMG128 / "face-chart.png", 6.4, "Graf posledných hodnôt, dole v kruhu."),
        ],
    )
    add_picture(
        doc,
        IMG128 / "face-source.png",
        6.4,
        "Keď je v menu zapnuté Zdroj, pod vekom je názov služby. Predvolene je vypnutý.",
    )

    add_heading_sk(doc, "Menu na displeji", 1)
    add_text(
        doc,
        "Menu má štyri strany. Medzi nimi sa prechádza potiahnutím doľava a doprava. "
        "Na kruhu sú tlačidlá pod sebou, takže Systém treba posunúť, aby bolo vidno spodok. "
        "Späť zatvorí menu.",
    )
    add_pair(
        doc,
        [
            (IMG128 / "menu-system.png", 6.2, "Systém 1/4, vrch: Informácie, Obnoviť, automatická aktualizácia."),
            (IMG128 / "menu-system-2.png", 6.2, "Ten istý Systém, posunutý: Aktualizácia, Zmazať WiFi, PIN, Späť."),
        ],
    )
    add_text(doc, "Čo tlačidlá na Systéme robia:", space_after=4)
    add_bullet(doc, "Informácie — IP adresa dosky, sila WiFi, teplota čipu a verzia firmvéru.")
    add_bullet(doc, "Obnoviť — hneď stiahne novú hodnotu.")
    add_bullet(doc, "Auto-upd — či sa má firmvér aktualizovať sám. Predvolene je zapnuté.")
    add_bullet(doc, "Aktualizácia — skontroluje a stiahne nový firmvér teraz.")
    add_bullet(doc, "Zmazať WiFi — zmaže úplne všetko: WiFi, účty aj PIN. Pýta sa znova.")
    add_bullet(doc, "PIN — nastaví alebo zmení štvormiestny PIN stránky. Keď PIN už je, pribudne Zrušiť PIN. To zmaže iba PIN, WiFi a účty ostanú.")

    add_picture(doc, IMG128 / "info.png", 6.2, "Informácie. IP na obrázku je ukážka.")
    add_pair(
        doc,
        [
            (IMG128 / "confirm-erase.png", 6.2, "Zmazať WiFi sa musí potvrdiť. Áno, zmazať vymaže všetko."),
            (IMG128 / "pin.png", 6.2, "PIN sú štyri číslice. Uložený je len odtlačok, nie samotné čísla."),
        ],
    )

    add_heading_sk(doc, "Limity, vzhľad a jas", 1)
    add_text(
        doc,
        "Limity sú spoločné pre obe osoby, keď ich nastavuješ z displeja. Predvolené sú spodný 3,9, "
        "vrchný 10,0 a kritické 15,0 mmol/l. Ťuknutím sa otvorí posuvník, Uložiť hodnotu zapíše.",
    )
    add_pair(
        doc,
        [
            (IMG128 / "menu-limits.png", 6.2, "Limity 2/4."),
            (IMG128 / "edit-low.png", 6.2, "Úprava spodného cieľa."),
        ],
    )
    add_text(doc, "Vzhľad 3/4:", space_after=4)
    add_bullet(doc, "Jednotky — mmol/l alebo mg/dl.")
    add_bullet(doc, "Delta — malé číslo zmeny vedľa šípky. Predvolene zapnuté.")
    add_bullet(doc, "Zdroj — názov služby pod číslom. Predvolene vypnutý.")
    add_bullet(doc, "Žiara — svietiaci kruh okolo čísla. Predvolene zapnutá.")
    add_bullet(doc, "Jazyk — prepínač v kruhu názvu jazyka, na obrázku Slovenčina.")
    add_picture(doc, IMG128 / "menu-look.png", 6.2, "Vzhľad. Na kruhu treba stranu posunúť, ak tlačidlo nie je vidno.")

    add_text(
        doc,
        "Jas a čas 4/4. Cez deň svieti displej naplno (255), v noci tlmene (25). "
        "Deň začína o 7:00, noc o 22:00. Hodiny sa dajú zmeniť posuvníkom.",
    )
    add_picture(doc, IMG128 / "menu-bright.png", 6.2, "Jas a začiatok dňa a noci.")

    add_heading_sk(doc, "Široká doska: dve osoby", 1)
    add_text(
        doc,
        "Na širokej doske má Systém navyše dve tlačidlá. Režim prepína medzi jednou osobou a dvoma vedľa seba. "
        "Zobraz prepína, ktorá osoba je na celej obrazovke, keď je režim jednej osoby a obe sú nastavené. "
        "Mená a účty sa dopĺňajú na stránke dosky, nie na tomto displeji.",
    )
    add_picture(
        doc,
        IMG35 / "menu-system.png",
        14.6,
        "Systém na šírku, dva stĺpce. Režim: 1 osoba a Zobraz: Anna sú len na tejto doske.",
    )
    add_text(
        doc,
        "V režime dvoch osôb sú vedľa seba dva stĺpce. Farba rámu je tá istá ako farba kruhu: "
        "zelená v cieli, oranžová nad vrchným limitom, červená pri spodnom limite a pri kritickej hodnote. "
        "Každý stĺpec má vlastné meno, číslo, šípku, zmenu, vek a zdroj. "
        "Potiahnutie doľava alebo doprava tu nestrieda hodiny. Strieda čísla a grafy. Hodiny ostanú len v režime jednej osoby.",
    )
    add_picture(
        doc,
        IMG35 / "face-split.png",
        14.6,
        "Dve osoby. Anna 6,4 zo zdroja Nightscout, Adam 12,4 zo zdroja Libre. Hodnoty sú ukážka.",
    )
    add_picture(
        doc,
        IMG35 / "face-split-chart.png",
        14.6,
        "Ten istý režim po potiahnutí: pod číslom je graf každej osoby.",
    )
    add_text(
        doc,
        "Keď sú nastavené dve osoby a necháš režim jednej, na ciferníku je len vybraná osoba a jej meno je vľavo od kruhu. "
        "To je obraz vyššie v úvode, s menom Anna.",
    )

    add_heading_sk(doc, "Široká doska: zvuk", 1)
    add_text(
        doc,
        "Reproduktor má len široká doska. Kruh zvuk nemá. Zvuk je predvolene vypnutý. "
        "Tlačidlá sú na strane Limity.",
    )
    add_picture(
        doc,
        IMG35 / "menu-limits.png",
        14.6,
        "Limity a zvuk. Snooze cyklí 15, 30 a 60 minút. Predvolene je 30. Stlmiť teraz stíši prebiehajúci alarm.",
    )
    add_text(
        doc,
        "Zvuk otvorí samostatnú stranu. Tam sa alarm zapne alebo vypne a nastaví hlasitosť: "
        "Tichá, Nízka, Stredná alebo Vysoká. Predvolená je Stredná.",
    )
    add_picture(doc, IMG35 / "menu-sound.png", 14.6, "Strana zvuku. Zvuk: VYP, hlasitosť Stredná.")
    add_text(
        doc,
        "Keď je zvuk zapnutý a niektorá hodnota je kritická (červená), krátke ťuknutie na ciferník alarm stíši "
        "na dobu z tlačidla Snooze. To isté spraví Stlmiť teraz.",
    )

    add_heading_sk(doc, "Aktualizácia firmvéru", 1)
    add_text(
        doc,
        "Doska sťahuje firmvér len cez HTTPS a len keď sedí podpis. Iný, nepodpísaný obraz odmietne. "
        "Prepínač TLS bez overenia na stránke platí pre Nightscout a podobné služby, nie pre túto aktualizáciu. "
        "Obraz kruhu a obraz širokej dosky sú dva rôzne súbory. Jeden do druhej dosky nepatrí.",
    )
    add_text(
        doc,
        "Počas sťahovania displej ukáže percentá. Nechaj dosku v dosahu WiFi, kým pruh nedobehne a doska sa reštartuje.",
    )
    add_picture(doc, IMG35 / "ota.png", 14.0, "Priebeh aktualizácie. 42 % je ukážka uprostred sťahovania.")

    add_heading_sk(doc, "Odkiaľ berie číslo", 1)
    add_text(
        doc,
        "Nightscout, Libre LinkUp, Dexcom Share a Guardian (CareLink) sú tu napojené neoficiálne. "
        "Nie sú to verejné služby tohto displeja. Výrobca ich môže zmeniť a podmienky služby ich nemusia dovoliť. "
        "GlukoView nie je schválená aplikácia Libre, Dexcom ani Medtronic.",
    )
    add_text(
        doc,
        "Dexcom sa prihlasuje údajmi z aplikácie u pacienta, nie účtom followera, a zdieľanie musí byť zapnuté. "
        "Dve osoby môžu mať každý vlastný účet, alebo pri Libre a CareLink jedno konto a rôzneho pacienta. "
        "To sa nastavuje na stránke dosky.",
    )
    add_text(
        doc,
        "Prázdne pole hesla pri uložení na stránke znamená, že pôvodné heslo ostáva. "
        "Nové heslo sa zapíše, len keď ho do poľa naozaj napíšeš.",
    )

    add_heading_sk(doc, "Čo si z toho vziať", 1)
    add_bullet(doc, "Číslo na displeji je orientačné. Dávku a liečbu ním samotným nerozhoduj.")
    add_bullet(doc, "Červená je spodný limit alebo kritická hodnota, oranžová je nad vrchným limitom, sivá so slovom Staré je meranie staršie ako 15 minút.")
    add_bullet(doc, "Menu otvorí podržanie prsta alebo krátke stlačenie BOOT.")
    add_bullet(doc, "Zmazať WiFi maže všetko. Zrušiť PIN maže len PIN.")
    add_bullet(doc, "Dve osoby vedľa seba a zvuk sú len na širokej doske.")

    doc.core_properties.title = "GlukoView — návod na použitie"
    doc.core_properties.subject = "Návod k stolnému ukazovateľu glykémie"
    doc.core_properties.language = "sk-SK"
    doc.core_properties.category = "Návod"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
