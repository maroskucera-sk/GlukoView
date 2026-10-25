# -*- coding: utf-8 -*-
"""Slovak Word guide: download GlukoView from GitHub and flash a board."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
IMG = ROOT / "docs" / "img" / "flash"
OUT = ROOT / "docs" / "GlukoView-stiahnutie-a-nahratie.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="C4A0A0", sz="8"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def add_text(doc, text, *, size=11, bold=False, space_after=8, space_before=0, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text) if not p.runs else p.runs[0]
    if p.runs and p.runs[0].text != text:
        p.runs[0].text = text
        run = p.runs[0]
    set_run_font(run)
    return p


def add_heading_sk(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=18 if level == 1 else 14, bold=True, color=RGBColor(0x14, 0x3D, 0x4A))
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_picture(doc, name, width_cm, caption):
    path = IMG / name
    if not path.is_file():
        raise SystemExit(f"chýba obrázok: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    set_run_font(
        cap.add_run(caption),
        size=9,
        italic=True,
        color=RGBColor(0x55, 0x55, 0x55),
    )


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, "F4F6F7")
    set_cell_border(cell, "D0D5D8", "4")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(2)
    for i, line in enumerate(text.split("\n")):
        if i:
            p.add_run("\n")
        set_run_font(p.add_run(line), name="Consolas", size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        footer.add_run("GlukoView  ·  stiahnutie a nahratie  ·  obraz z jednej dosky nepatrí do druhej"),
        size=8,
        color=RGBColor(0x88, 0x88, 0x88),
    )

    add_text(doc, "GlukoView", size=26, bold=True, space_after=0, color=RGBColor(0x14, 0x3D, 0x4A))
    add_text(doc, "Stiahnutie z GitHubu a nahratie do dosky", size=16, space_after=8, color=RGBColor(0x3A, 0x7A, 0x4A))
    add_text(
        doc,
        "Tento návod je na to, ako dostať program z GitHubu do dosky cez USB. "
        "Ako sa potom číselník ovláda, je v súbore GlukoView-navod.docx. "
        "Číslo, ktoré doska po nahratí ukáže, je stále len orientačné a nie je zdravotnícka pomôcka.",
    )

    add_heading_sk(doc, "Čo treba", 1)
    add_bullet(doc, "Počítač s Windows 10 alebo 11.")
    add_bullet(doc, "Kábel USB, ktorý prenáša dáta. Kábel len na nabíjanie dosku v počítači neukáže.")
    add_bullet(doc, "Prehliadač, Visual Studio Code a doň doplnok PlatformIO IDE.")
    add_bullet(doc, "Internet. Pri prvom preklade si PlatformIO stiahne prekladač a knižnice. Trvá to dlhšie.")
    add_bullet(doc, "Git len vtedy, keď namiesto ZIPu chceš príkaz git clone. Na Download ZIP netreba.")

    add_heading_sk(doc, "1. Stiahnuť kód z GitHubu", 1)
    add_text(
        doc,
        "Keď je kód na GitHube, otvor odkaz na repozitár GlukoView. Vpravo je zelené tlačidlo Code. "
        "Obrázok ukazuje to isté tlačidlo na inom, verejnom repozitári, aby bolo vidno menu. "
        "Názov platformio-core na obrázku nie je tento projekt. Ty otvoríš svoj odkaz. Menu vyzerá rovnako.",
    )
    add_picture(
        doc,
        "github-code.png",
        15.2,
        "Zelené Code. Download ZIP stiahne priečinok. Adresa pod HTTPS je na git clone.",
    )
    add_text(doc, "Jednoduchšia cesta, bez Gitu:", space_after=4)
    add_bullet(doc, "Code → Download ZIP.")
    add_bullet(doc, "ZIP rozbaľ. Vnútri musí byť súbor platformio.ini. Rozbalený priečinok sa často volá s koncovkou -main. Otvor ten, v ktorom platformio.ini naozaj je, nie priečinok o úroveň vyššie.")
    add_text(doc, "Druhá cesta, ak máš nainštalovaný Git:", space_after=4)
    add_bullet(doc, "V menu Code ostaneš na HTTPS a skopíruješ adresu.")
    add_bullet(doc, "Vo PowerShelli, v priečinku, kam chceš kód:")
    add_code(doc, "git clone ADRESA-Z-GITHUBU")
    add_text(
        doc,
        "Ak príkaz git počítač nepozná, nainštaluj Git z https://git-scm.com/download/win a nechaj ponúknuté voľby. "
        "Potom príkaz zopakuj.",
    )

    add_heading_sk(doc, "2. Visual Studio Code", 1)
    add_text(
        doc,
        "Program, v ktorom sa projekt otvorí a nahrá, je Visual Studio Code. Stránka sama ponúkne Windows. "
        "Stiahni Windows a inštalátor nechaj v predvolených krokoch.",
    )
    add_picture(
        doc,
        "vscode-download.png",
        15.2,
        "code.visualstudio.com/download — tlačidlo Windows. Na tomto počítači je označené ako tvoj systém.",
    )

    add_heading_sk(doc, "3. Doplnok PlatformIO", 1)
    add_text(
        doc,
        "Samotné Visual Studio Code dosku nenahrá. Treba doplnok PlatformIO IDE. "
        "Vo Visual Studio Code klikni vľavo na ikonu štvorcov (Extensions), do hľadania napíš PlatformIO IDE a daj Install. "
        "To isté Install je na stránke doplnku. Po inštalácii Visual Studio Code raz reštartuj, keď si to vypýta.",
    )
    add_picture(
        doc,
        "platformio-install.png",
        15.2,
        "Doplnok PlatformIO IDE. Zelené Install. V obchode Visual Studio Code sa hľadá tým istým menom.",
    )
    add_text(
        doc,
        "Pri prvom otvorení projektu si PlatformIO dotiahne nástroje Espressif. Nechaj to dobehnúť a neodpájaj internet. "
        "Môže to trvať aj desiatky minút. Ďalšie nahratia sú už kratšie.",
    )

    add_heading_sk(doc, "4. Otvoriť projekt", 1)
    add_bullet(doc, "File → Open Folder.")
    add_bullet(doc, "Vyber priečinok, v ktorom leží platformio.ini.")
    add_bullet(doc, "Ak sa opýta, či autorom veríš, potvrď Trust. Inak sa projekt nepreloží.")

    add_heading_sk(doc, "5. Vybrať dosku", 1)
    add_text(
        doc,
        "Vľavo je ikona PlatformIO, hlava mimozemšťana. Pod Project Tasks je General a v ňom Upload. "
        "Upload nahrá tú dosku, ktorej meno je na modrom pruhu dole. Na obrázku je to env:waveshare_s3_3_5, čiže široká doska. "
        "Klikni na to meno a vyber prostredie podľa dosky, ktorú máš v ruke.",
    )
    add_bullet(doc, "env:waveshare_s3_1_28 — kruhový Waveshare 1,28″.")
    add_bullet(doc, "env:waveshare_s3_3_5 — široký Waveshare ESP32-S3-Touch-LCD-3.5B.")
    add_bullet(doc, "env:waveshare_s3_3_5_manual — len keď bežné nahratie širokej dosky nechytí bootloader. Postup je nižšie.")
    add_text(
        doc,
        "Položku Default v tom zozname neber. Projekt nemá nastavené default_envs, takže Default nie je jedna konkrétna doska. "
        "Súbor platformio.ini v editore ukazuje obe dosky pod sebou. Platí meno na modrom pruhu, nie odsek, na ktorý sa práve pozeráš.",
    )
    add_picture(
        doc,
        "vscode-upload.png",
        16.0,
        "PlatformIO vľavo, Upload pod General. Dole na modrom pruhu je vybraná doska env:waveshare_s3_3_5.",
    )

    add_heading_sk(doc, "6. Pripojiť dosku", 1)
    add_text(
        doc,
        "Dosku pripoj dátovým káblom. Vo Visual Studio Code otvor Terminal → New Terminal a spusti:",
    )
    add_code(doc, "pio device list")
    add_text(
        doc,
        "Doska tohto projektu je riadok USB Serial Device, ktorého Hardware ID začína USB VID:PID=303A. "
        "303A je Espressif, čip ESP32-S3, a nahráva sa cez jeho vlastné USB. "
        "Na obrázku je to COM4. Číslo COM je na každom počítači iné. Hľadáš 303A, nie konkrétne číslo.",
    )
    add_picture(
        doc,
        "pio-ports.png",
        14.5,
        "pio device list. Doska je COM4, USB Serial Device, VID 303A. CH343 a NPort sú iné zariadenia.",
    )
    add_bullet(doc, "CH343 (VID 1A86) je iný prevodník. Pre toto nahratie ho nevyberaj, keď je v zozname aj 303A.")
    add_bullet(doc, "COM1 typu NPort nie je doska.")
    add_bullet(doc, "Keď 303A v zozname nie je, kábel je skoro isto len nabíjací, alebo je v inom porte. Skús iný kábel a iný USB port. Ovládač na Windows 10 a 11 príde sám. CP210x ani CH340 sa pre túto dosku neinštalujú.")

    add_heading_sk(doc, "7. Nahrať", 1)
    add_text(
        doc,
        "Keď je na modrom pruhu správne prostredie, klikni Project Tasks → General → Upload. "
        "Istota, ktorá nezávisí od modrého pruhu, je príkaz v tom istom termináli. Pre kruh:",
    )
    add_code(doc, "pio run -t upload -e waveshare_s3_1_28")
    add_text(doc, "Pre širokú dosku:", space_before=4)
    add_code(doc, "pio run -t upload -e waveshare_s3_3_5")
    add_text(
        doc,
        "Prvý beh sťahuje platformu espressif32 6.3.1 a knižnice. Nechaj ho dobehnúť a dosku neodpájaj. "
        "Na konci má byť riadok SUCCESS. Doska sa reštartuje a ukáže logo. Číslo verzie na logu je to, "
        "ktoré je v platformio.ini v riadkoch FIRMWARE_VERSION. Bežné nahratie nezmaže uložené WiFi ani účty.",
    )
    add_text(
        doc,
        "Ak pred ďalším nahratím beží sledovanie portu, najprv ho zatvor. Inak je port obsadený a upload spadne. "
        "Sledovanie, keď ho chceš, je:",
    )
    add_code(doc, "pio device monitor -b 115200")

    add_heading_sk(doc, "8. Keď široká doska nenaskočí", 1)
    add_text(
        doc,
        "Pri širokej doske sa občas bežný upload nestihne chytiť. Vtedy ju daj do režimu nahrávania ručne: "
        "podrž BOOT, stlač RST, pusť BOOT. Potom, kým je stále pripojená:",
    )
    add_code(doc, "pio run -t upload -e waveshare_s3_3_5_manual")
    add_text(
        doc,
        "To isté prostredie je v zozname na modrom pruhu ako env:waveshare_s3_3_5_manual a potom Upload. "
        "Toto prostredie dosku pred zápisom samo neresetuje, preto musia byť tlačidlá stlačené skôr. "
        "Kruhová doska toto ručné prostredie nemá. Pri nej stačí waveshare_s3_1_28.",
    )

    add_heading_sk(doc, "Čo z GitHubu nepríde hotové", 1)
    add_text(
        doc,
        "Z GitHubu stiahneš zdrojový kód, nie hotový súbor na prekopírovanie do dosky. "
        "Akcie na GitHube tie dva obrazy preložia, ale firmware.bin na stiahnutie nezverejňujú. "
        "Súbor vznikne až prekladom u teba, v priečinku:",
    )
    add_code(
        doc,
        ".pio\\build\\waveshare_s3_1_28\\firmware.bin\n"
        ".pio\\build\\waveshare_s3_3_5\\firmware.bin",
    )
    add_text(
        doc,
        "Pri USB nahratí ho nikam nekopíruješ. PlatformIO ho do dosky zapíše samo. "
        "Obraz kruhu do širokej dosky nepatrí a obraz širokej dosky nepatrí do kruhu.",
    )

    doc.core_properties.title = "GlukoView — stiahnutie a nahratie"
    doc.core_properties.subject = "Ako stiahnuť kód z GitHubu a nahrať ho do dosky"
    doc.core_properties.language = "sk-SK"
    doc.core_properties.category = "Návod"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
